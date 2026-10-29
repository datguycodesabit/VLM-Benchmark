"""Product-level checks across generated configuration, execution, and offline reports."""

import csv
import json

import pytest
from openpyxl import load_workbook
from PIL import Image
from test_engine import Backend

from vlm_bench import cli, config, engine
from vlm_bench.init import create_config
from vlm_bench.reporting import write_report
from vlm_bench.snapshot import freeze


def _data(tmp_path):
    data = tmp_path / "input"
    (data / "images").mkdir(parents=True)
    (data / "text").mkdir()
    for index in range(3):
        Image.new("RGB", (40 + index, 15), "white").save(data / "images" / f"s{index}.png")
        (data / "text" / f"s{index}.txt").write_text("hello")
    return data


def test_init_roundtrips_and_preserves_existing_config(tmp_path):
    path = create_config(
        tmp_path / "benchmark.toml", prepared=tmp_path / "prepared", models=["ollama:fake"]
    )
    parsed = config.load_config(path)
    assert parsed["version"] == 3
    assert parsed["suite"]["conditions"][0]["repeats"] == 3
    assert parsed["suite"]["conditions"][0]["prepared"] == tmp_path / "prepared"
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        create_config(path, prepared=tmp_path / "other", models=["ollama:other"])
    assert path.read_bytes() == before


def test_direct_cli_three_passes_and_shared_offline_exports(tmp_path, monkeypatch, capsys):
    data = _data(tmp_path)
    original = engine.run
    Backend.calls = []
    Backend.failure = None
    Backend.identity = "fixed"
    monkeypatch.setattr(engine, "run", lambda *a, **kw: original(*a, **kw, backend_factory=Backend))
    output = tmp_path / "runs"
    assert (
        cli.main(
            [
                "run",
                "--data",
                str(data),
                "--models",
                "ollama:fake",
                "--output",
                str(output),
                "--no-warmup",
            ]
        )
        == 0
    )
    capsys.readouterr()
    experiment = next(output.iterdir())
    state = json.loads((experiment / "experiment.json").read_text())
    assert len(state["runs"]) == state["planned_passes"] == 3
    assert len(Backend.calls) == 9
    manifests = [
        json.loads((experiment / path / "manifest.json").read_text()) for path in state["runs"]
    ]
    assert [value["repetition"] for value in manifests] == [1, 2, 3]
    assert len({value["benchmark_fingerprint"] for value in manifests}) == 1
    assert all(value["planned_passes"] == 3 for value in manifests)
    report = json.loads((experiment / "report" / "evaluation.json").read_text())
    rows = report["conditions"]["baseline"]["tracks"]["prose"]["models"]
    assert rows[0]["score"] == 100
    assert report["conditions"]["baseline"]["finalizable"]
    csv_rows = list(
        csv.DictReader((experiment / "report" / "leaderboards.csv").open(encoding="utf-8-sig"))
    )
    assert float(csv_rows[0]["score"]) == rows[0]["score"]
    workbook = load_workbook(experiment / "report" / "report.xlsx", read_only=True)
    values = list(workbook["Leaderboards"].values)
    assert values[1][values[0].index("score")] == rows[0]["score"]
    workbook.close()
    assert (experiment / "report" / "report.html").is_file()
    assert list((experiment / "report" / "assets").iterdir())
    monkeypatch.setattr(engine, "run", lambda *a, **kw: pytest.fail("offline report invoked model"))
    write_report(experiment, tmp_path / "regenerated")
    assert len(Backend.calls) == 9


def test_html_escapes_predictions_and_excel_protects_formula_text(tmp_path):
    data = _data(tmp_path)
    prepared = tmp_path / "prepared"
    freeze(data, prepared)
    from vlm_bench.import_predictions import import_predictions

    path = tmp_path / "predictions.jsonl"
    path.write_text(
        "\n".join(
            json.dumps(
                {
                    "sample_id": f"s{i}",
                    "prediction": "<script>alert(1)</script>" if i == 0 else "=1+1",
                }
            )
            for i in range(3)
        )
    )
    run = import_predictions(
        prepared, path, tmp_path / "imported", system="test", provenance="synthetic export fixture"
    )
    write_report(run)
    text = (run / "report" / "report.html").read_text()
    assert "<script>alert(1)</script>" not in text
    assert "&lt;script&gt;" in text
    workbook = load_workbook(run / "report" / "report.xlsx")
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_generated_v3_suite_runs_and_resumes_without_remeasurement(tmp_path, monkeypatch, capsys):
    data = _data(tmp_path)
    prepared = tmp_path / "prepared"
    freeze(data, prepared)
    cfg = create_config(tmp_path / "benchmark.toml", prepared=prepared, models=["ollama:fake"])
    with cfg.open("a") as stream:
        stream.write("\n[scoring]\ncharacter_accuracy_weight=0.6\nexact_match_weight=0.4\n")
    original_run, original_preview = engine.run, engine.preview
    Backend.calls, Backend.failure, Backend.identity = [], None, "fixed"
    monkeypatch.setattr(
        engine, "run", lambda *a, **kw: original_run(*a, **kw, backend_factory=Backend)
    )
    monkeypatch.setattr(
        engine, "preview", lambda *a, **kw: original_preview(*a, **kw, backend_factory=Backend)
    )
    output = tmp_path / "suite-output"
    assert cli.main(["suite", "--config", str(cfg), "--output", str(output)]) == 0
    capsys.readouterr()
    calls = len(Backend.calls)
    assert calls == 12  # Three measured passes and their separate warmups.
    assert (
        cli.main(["suite", "--config", str(cfg), "--output", str(output), "--resume", "--json"])
        == 0
    )
    payload = json.loads(capsys.readouterr().out)
    assert len(Backend.calls) == calls
    assert any(path.endswith("report.html") for path in payload["report_paths"])
    report = json.loads((output / "report" / "evaluation.json").read_text())
    assert report["conditions"]["baseline"]["policy"]["exact_match_weight"] == 0.4
    assert report["conditions"]["baseline"]["finalizable"]


def test_report_workbook_preserves_long_text_in_overflow_sheet(tmp_path, monkeypatch):
    from vlm_bench import report_model

    long_text = "=custom-text-" + "x" * 40000
    monkeypatch.setattr(report_model, "build_report", lambda *a, **kw: {"notes": long_text})
    write_report(tmp_path, tmp_path / "exported", formats=["xlsx"])
    workbook = load_workbook(tmp_path / "exported" / "report.xlsx")
    chunks = list(workbook["Text Overflow"].values)[1:]
    assert "".join(row[-1] for row in chunks) == long_text
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
