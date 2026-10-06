from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from vlm_bench.export import _EXCEL_JSON_STRING_PREFIX, export_run
from vlm_bench.metrics import score


def _write_run(
    run_dir: Path,
    records: list[dict],
    manifest: dict,
    warmups: list[dict] | None = None,
) -> bytes:
    run_dir.mkdir(parents=True, exist_ok=True)
    result_bytes = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in records).encode(
        "utf-8"
    )
    (run_dir / "results.jsonl").write_bytes(result_bytes)
    (run_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False), encoding="utf-8"
    )
    if warmups is not None:
        (run_dir / "warmups.jsonl").write_text(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in warmups),
            encoding="utf-8",
        )
    return result_bytes


def _success_record(model: str, sample_id: str, prediction: str, reference: str) -> dict:
    return {
        "model": model,
        "sample_id": sample_id,
        "status": "success",
        "prediction": prediction,
        "reference": reference,
        "metrics": score(prediction, reference),
        "latency_seconds": 1.25,
        "load_duration_seconds": 0.1,
        "truncated": False,
        "error": None,
    }


def test_export_run_writes_summary_csv_samples_csv_and_xlsx_without_mutating_jsonl(tmp_path):
    formula_text = "=1+1"
    records = [
        _success_record(formula_text, "s1", formula_text, formula_text),
        {
            "model": "broken",
            "sample_id": "s1",
            "status": "error",
            "prediction": None,
            "reference": "a",
            "metrics": None,
            "latency_seconds": 4.0,
            "load_duration_seconds": None,
            "truncated": False,
            "error": "=TIMEOUT()",
        },
    ]
    records[0].update(
        cache_hit=True,
        cache_bypass_reason=None,
        cache_lookup_seconds=0.003,
        latency_seconds=None,
        inference_latency_seconds=None,
    )
    records[1].update(cache_hit=False, cache_bypass_reason="cache_disabled")
    warmups = [
        {
            "model": formula_text,
            "sample_id": "s1",
            "status": "success",
            "latency_seconds": 2.0,
            "response": {"load_duration": 1_500_000_000},
        }
    ]
    original_jsonl = _write_run(
        tmp_path,
        records,
        {
            "status": "complete_with_errors",
            "models": [formula_text, "broken"],
            "samples": [{"id": "s1"}],
            "options": {"temperature": 0, "seed": 42},
        },
        warmups,
    )

    outputs = export_run(tmp_path, ["xlsx", "csv", "jsonl"])

    assert outputs == [
        tmp_path / "results.xlsx",
        tmp_path / "summary.csv",
        tmp_path / "samples.csv",
        tmp_path / "results.jsonl",
    ]
    assert (tmp_path / "results.jsonl").read_bytes() == original_jsonl

    with (tmp_path / "samples.csv").open(encoding="utf-8-sig", newline="") as file:
        samples = {row["model"]: row for row in csv.DictReader(file)}
    assert samples["'=1+1"]["prediction"] == "'=1+1"
    assert samples["broken"]["prediction"] == ""
    assert samples["broken"]["error"] == "'=TIMEOUT()"
    assert samples["'=1+1"]["cer"] == "0.0"
    assert samples["'=1+1"]["cache_hit"] == "true"
    assert samples["'=1+1"]["cache_bypass_reason"] == ""
    assert samples["'=1+1"]["cache_lookup_seconds"] == "0.003"
    assert samples["broken"]["cache_bypass_reason"] == "cache_disabled"

    with (tmp_path / "summary.csv").open(encoding="utf-8-sig", newline="") as file:
        summaries = {row["model"]: row for row in csv.DictReader(file)}
    assert summaries["'=1+1"]["complete"] == "true"
    assert summaries["'=1+1"]["warmup_count"] == "1"
    assert summaries["'=1+1"]["warmup_load_duration_seconds"] == "1.5"
    assert summaries["'=1+1"]["cache_hit_count"] == "1"
    assert summaries["'=1+1"]["measured_latency_sample_count"] == "0"
    assert summaries["broken"]["complete"] == "false"
    assert summaries["broken"]["failure_rate"] == "1.0"
    assert summaries["broken"]["cache_hit_count"] == "0"
    assert summaries["broken"]["measured_latency_sample_count"] == "1"

    workbook = load_workbook(tmp_path / "results.xlsx", data_only=False)
    assert workbook.sheetnames == ["Summary", "Samples", "Errors", "Run Configuration", "Warmups"]
    summary_headers = {cell.value for cell in workbook["Summary"][1]}
    assert {"cache_hit_count", "measured_latency_sample_count"}.issubset(summary_headers)
    sample_headers = {cell.value for cell in workbook["Samples"][1]}
    assert {"cache_hit", "cache_bypass_reason"}.issubset(sample_headers)
    assert not [
        cell
        for sheet in workbook.worksheets
        for row in sheet
        for cell in row
        if cell.data_type == "f"
    ]

    samples_sheet = workbook["Samples"]
    headers = {cell.value: cell.column for cell in samples_sheet[1]}
    formula_model_row = next(
        row
        for row in range(2, samples_sheet.max_row + 1)
        if samples_sheet.cell(row, headers["model"]).value == formula_text
    )
    model_cell = samples_sheet.cell(formula_model_row, headers["model"])
    prediction_cell = samples_sheet.cell(formula_model_row, headers["prediction"])
    assert model_cell.value == formula_text and model_cell.data_type == "s"
    assert prediction_cell.value == formula_text and prediction_cell.data_type == "s"

    assert workbook["Errors"].max_row == 2
    config = {
        workbook["Run Configuration"].cell(row, 1).value: workbook["Run Configuration"]
        .cell(row, 2)
        .value
        for row in range(2, workbook["Run Configuration"].max_row + 1)
    }
    assert config["options.temperature"] == "0"
    assert workbook["Warmups"].max_row == 2


def test_xlsx_overflow_and_control_encoding_preserve_text_recovery(tmp_path):
    long_text = "🙂" * 17_000 + "tail"
    control_text = "prefix\x01middle\x1fend"
    metrics = {
        "cer": 0.0,
        "wer": 0.0,
        "raw_cer": 0.0,
        "exact_match": True,
        "char_substitutions": 0,
        "char_deletions": 0,
        "char_insertions": 0,
        "char_edits": 0,
        "word_edits": 0,
        "reference_chars": 1,
        "reference_words": 1,
    }
    records = [
        {
            "model": "vision",
            "sample_id": "long",
            "status": "success",
            "prediction": long_text,
            "reference": "x",
            "metrics": metrics,
        },
        {
            "model": "vision",
            "sample_id": "control",
            "status": "success",
            "prediction": control_text,
            "reference": "x",
            "metrics": metrics,
        },
    ]
    _write_run(
        tmp_path,
        records,
        {
            "status": "complete",
            "models": ["vision"],
            "samples": [{"id": "long"}, {"id": "control"}],
        },
    )

    export_run(tmp_path, ["xlsx"])

    workbook = load_workbook(tmp_path / "results.xlsx", data_only=False)
    assert "Text Overflow" in workbook.sheetnames
    samples_sheet = workbook["Samples"]
    headers = {cell.value: cell.column for cell in samples_sheet[1]}
    rows_by_id = {
        samples_sheet.cell(row, headers["sample_id"]).value: row
        for row in range(2, samples_sheet.max_row + 1)
    }
    long_marker = samples_sheet.cell(rows_by_id["long"], headers["prediction"]).value
    assert long_marker.startswith("[FULL TEXT IN Text Overflow: ")
    overflow_sheet = workbook["Text Overflow"]
    overflow_headers = {cell.value: cell.column for cell in overflow_sheet[1]}
    overflow_id = long_marker.split(": ", 1)[1].rstrip("]")
    chunks = [
        overflow_sheet.cell(row, overflow_headers["text_chunk"]).value
        for row in range(2, overflow_sheet.max_row + 1)
        if overflow_sheet.cell(row, overflow_headers["overflow_id"]).value == overflow_id
    ]
    assert "".join(chunks) == long_text

    control_cell = samples_sheet.cell(rows_by_id["control"], headers["prediction"])
    assert control_cell.value.startswith(_EXCEL_JSON_STRING_PREFIX)
    assert json.loads(control_cell.value[len(_EXCEL_JSON_STRING_PREFIX) :]) == control_text


def test_interrupted_run_and_missing_manifest_samples_are_not_ranked(tmp_path):
    record = _success_record("vision", "one", "same", "same")
    _write_run(
        tmp_path,
        [record],
        {"status": "interrupted", "models": ["vision"], "samples": [{"id": "one"}, {"id": "two"}]},
    )

    export_run(tmp_path, ["csv"])

    with (tmp_path / "summary.csv").open(encoding="utf-8-sig", newline="") as file:
        summary = next(csv.DictReader(file))
    assert summary["complete"] == "false"
    assert summary["rank"] == ""
    assert summary["expected_sample_count"] == "2"
    assert summary["missing_sample_count"] == "1"


def test_research_export_includes_evaluation_protocol(tmp_path):
    record = _success_record("vision", "one", "same", "same")
    manifest = {
        "schema_version": 2,
        "status": "complete",
        "models": ["vision"],
        "samples": [
            {
                "id": "one",
                "reference": "same",
                "split": "test",
                "source_document": "doc-a",
                "verification_status": "verified",
            }
        ],
        "strict_research": True,
        "protocol": "writer-disjoint",
        "research_protocol_version": 1,
    }
    _write_run(tmp_path, [record], manifest)

    export_run(tmp_path, ["jsonl"])

    report = json.loads((tmp_path / "research.json").read_text(encoding="utf-8"))
    assert report["evaluation"] == {
        "strict_research": True,
        "protocol": "writer-disjoint",
        "research_protocol_version": 1,
    }


def test_schema_two_export_adds_unranked_task_metrics_sheet_and_csv(tmp_path):
    from vlm_bench.task_metrics import score_task

    sample = {
        "id": "equation-1",
        "reference": "x = 1",
        "hashes": {"crop": hashlib.sha256(b"crop").hexdigest()},
        "preprocess": "original",
        "content_type": "equation",
        "split": "test",
        "verified": True,
        "source_document": "doc-1",
        "metadata": {
            "annotations": {
                "critical_expressions": ["x"],
                "reading_order": [],
            }
        },
    }
    record = _success_record("vision", sample["id"], "x = 1", sample["reference"])
    record["metrics"]["task_metrics"] = score_task(
        record["prediction"], sample["reference"], sample
    )
    _write_run(
        tmp_path,
        [record],
        {
            "schema_version": 2,
            "status": "complete",
            "models": ["vision"],
            "samples": [sample],
            "preprocess": "original",
            "formula_rendering": False,
            "task_metrics_version": 1,
        },
    )

    outputs = export_run(tmp_path, ["xlsx", "csv"])
    report = json.loads((tmp_path / "research.json").read_text(encoding="utf-8"))
    assert report["task_metrics"]["ranking"] is None
    assert tmp_path / "task_metrics.csv" in outputs
    with (tmp_path / "task_metrics.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert any(row["track"] == "equation" for row in rows)

    workbook = load_workbook(tmp_path / "results.xlsx", data_only=True)
    assert "Task Metrics" in workbook.sheetnames
    assert "track" in {cell.value for cell in workbook["Task Metrics"][1]}
