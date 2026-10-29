"""Portable HTML, Excel, CSV and JSON views of one offline evaluation model."""

from __future__ import annotations

import csv
import hashlib
import html
import json
import math
import shutil
import statistics
from collections import defaultdict
from pathlib import Path

from .export import _append_table, _csv_value
from .runner import _atomic_json

_STYLE = """
:root {color-scheme:light;--ink:#172d40;--muted:#526879;--accent:#007b78}
*{box-sizing:border-box}body{margin:0;background:#f2f5f7;color:var(--ink);font:15px/1.6 system-ui,sans-serif}
main{max-width:1300px;margin:auto;padding:42px 24px}header{border-top:5px solid var(--accent);padding:25px 0}
h1{font-size:36px;line-height:1.15;margin:8px 0}h2{margin:0 0 16px}h3{text-transform:capitalize}
.eyebrow{color:var(--accent);font-weight:750;letter-spacing:.14em;text-transform:uppercase}
section{background:white;border:1px solid #dce5e9;border-radius:12px;padding:24px;margin:20px 0;overflow:auto}
p{max-width:95ch}.muted{color:var(--muted)}table{border-collapse:collapse;width:100%;font-size:13px}
th,td{text-align:left;padding:12px;border-bottom:1px solid #e2e9ed;vertical-align:top;max-width:420px;overflow-wrap:anywhere}
th{background:#eaf4f3;color:#1a5551}tr:hover{background:#f4f9f9}.num{font-variant-numeric:tabular-nums}
pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;background:#f4f7f8;padding:16px;border-radius:8px}
.table-scroll{overflow-x:auto}table{min-width:650px}th{min-width:90px}th:first-child,td:first-child{min-width:140px}.num{white-space:nowrap}
img{max-width:100%;max-height:180px;object-fit:contain}details{margin:12px 0}summary{cursor:pointer;font-weight:650}
@media print{body{background:white}main{padding:0}section{break-inside:avoid}}
"""


def _text(value):
    if value is None:
        return "Unavailable"
    if isinstance(value, float):
        return f"{value:.6g}" if math.isfinite(value) else "Unavailable"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return str(value)


def _escape(value):
    return html.escape(_text(value), quote=True)


def _evaluations(report):
    return report.get("conditions", {"Benchmark": report}).items()


def _leaderboards(report):
    rows = []
    for condition, evaluation in _evaluations(report):
        for track, data in evaluation.get("tracks", {}).items():
            results = data.get("models", data.get("model_results", []))
            if isinstance(results, dict):
                results = [dict(value, model=key) for key, value in results.items()]
            for row in results:
                rows.append({"condition": condition, "track": track, **row})
    return rows


def _flat(row):
    return {
        key: _text(value) if isinstance(value, (dict, list)) else value
        for key, value in row.items()
    }


def _table(rows):
    if not rows:
        return '<p class="muted">No eligible samples in this track.</p>'
    fields = list(dict.fromkeys(key for row in rows for key in row))
    return (
        '<div class="table-scroll"><table><thead><tr>'
        + "".join(f"<th>{_escape(key.replace('_', ' '))}</th>" for key in fields)
        + "</tr></thead><tbody>"
        + "".join(
            "<tr>"
            + "".join(
                f'<td class="num">{_escape(row.get(key))}</td>'
                if isinstance(row.get(key), (float, int))
                else f"<td>{_escape(row.get(key))}</td>"
                for key in fields
            )
            + "</tr>"
            for row in rows
        )
        + "</tbody></table></div>"
    )


def _html(report):
    parts = [
        '<!doctype html><html lang="en"><meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width,initial-scale=1">',
        "<title>VLM Benchmark · Research report</title>",
        f"<style>{_STYLE}</style>",
        '<main><header><div class="eyebrow">VLM Benchmark / Evaluation</div>',
        "<h1>Handwriting recognition benchmark</h1>",
        '<p class="muted">Separate prose and equation fidelity. Reproducible evidence, '
        "accuracy bands, and measured request latency.</p></header>",
    ]
    rows = _leaderboards(report)
    for name, evaluation in _evaluations(report):
        parts.append(f"<section><h2>{_escape(name)}</h2>")
        if not evaluation.get("tracks"):
            parts.append(f"<p>{_escape(evaluation.get('reason') or evaluation.get('status'))}</p>")
        policy = evaluation.get("policy", evaluation.get("scoring_policy", {}))
        parts.append(
            "<p>Score = 100 × (character weight × max(0, 1 − corpus CER) + "
            "exact-match weight × exact-match rate). Missing and failed outputs "
            "receive deletion penalties. Scores are averaged across planned passes.</p>"
        )
        parts.append(
            f"<details><summary>Frozen scoring policy</summary><pre>{_escape(policy)}</pre></details>"
        )
        for track, data in evaluation.get("tracks", {}).items():
            parts.append(f"<h3>{_escape(track)} leaderboard</h3>")
            ranking = data.get("ranking", data)
            parts.append(
                f"<p>Winner: <strong>{_escape(ranking.get('winner'))}</strong>. "
                f"{_escape(ranking.get('winner_reason') or ranking.get('winner_status', ''))}</p>"
            )
            candidates = [row for row in rows if row["condition"] == name and row["track"] == track]
            # Keep the principal comparison readable; full diagnostics follow separately.
            primary = [
                key
                for key in (
                    "model",
                    "placement",
                    "rank",
                    "accuracy_rank",
                    "raw_accuracy_rank",
                    "score",
                    "score_stddev",
                    "median_latency_seconds",
                    "complete",
                )
                if any(key in row for row in candidates)
            ]
            if not primary:
                primary = (
                    [key for key in candidates[0] if key not in {"condition", "track"}]
                    if candidates
                    else []
                )
            parts.append(_table([{key: row.get(key) for key in primary} for row in candidates]))
            measured = [
                row
                for row in candidates
                if isinstance(row.get("score"), (float, int))
                and isinstance(row.get("median_latency_seconds"), (float, int))
            ]
            if measured:
                maximum = max(row["median_latency_seconds"] for row in measured) or 1
                parts.append(
                    '<h4>Accuracy versus latency</h4><svg role="img" '
                    'aria-label="Accuracy score versus median latency" viewBox="0 0 700 250" '
                    'style="max-width:900px;width:100%;background:#f5f9fa">'
                    '<path d="M50 20V210H660" fill="none" stroke="#526879"/>'
                    '<text x="300" y="245" font-size="12">Median end-to-end latency (seconds) →</text>'
                    '<text x="10" y="18" font-size="11">Score</text>'
                )
                colors = ["#007b78", "#c37428", "#6253a0", "#ad4168"]
                for index, row in enumerate(measured):
                    x = 60 + row["median_latency_seconds"] / maximum * 550
                    y = 210 - row["score"] * 1.8
                    parts.append(
                        f'<circle cx="{x:.2f}" cy="{y:.2f}" r="6" fill="{colors[index % 4]}" opacity=".8">'
                        f"<title>{_escape(row['model'])}: {_escape(row['score'])} / "
                        f"{_escape(row['median_latency_seconds'])} seconds</title></circle>"
                    )
                parts.append(
                    f'<text x="14" y="34" font-size="11">100</text>'
                    f'<text x="32" y="212" font-size="11">0</text>'
                    f'<text x="600" y="226" font-size="11">{_escape(maximum)}</text></svg>'
                )
                parts.append(
                    '<p class="muted">'
                    + " · ".join(
                        f'<span style="color:{colors[index % 4]}">● {_escape(row["model"])}</span>'
                        for index, row in enumerate(measured)
                    )
                    + "</p>"
                )
            passes = [
                {"model": row["model"], **item}
                for row in candidates
                for item in row.get("per_pass", [])
                if isinstance(item, dict)
            ]
            if passes:
                parts.append(
                    "<details><summary>Per-pass measurements</summary>"
                    + _table(passes)
                    + "</details>"
                )
            parts.append(
                "<details><summary>Coverage, per-pass scores, uncertainty and equivalence</summary>"
            )
            parts.append(
                f"<pre>{_escape(json.dumps(data, ensure_ascii=False, indent=2))}</pre></details>"
            )
        parts.append("</section>")
    parts.append(
        "<section><h2>Reproducibility and measurement context</h2><p>Latency uses "
        "successful uncached sequential requests. Unavailable timing and cost remain "
        "unknown. Equation equivalence is diagnostic and does not change fidelity rankings. "
        "Hardware and provider differences are part of each measured system.</p>"
    )
    parts.append("<h3>Operational evidence</h3>" + _table(report.get("operational_summary", [])))
    parts.append("<h3>Available timing breakdowns</h3>")
    if report.get("timing_breakdowns"):
        parts.append(_table(report["timing_breakdowns"]))
    else:
        parts.append(
            '<p class="muted">No instrumented breakdowns were recorded. '
            "Cloud internals and time to first output remain unavailable.</p>"
        )
    context = {
        key: value
        for key, value in report.items()
        if key not in {"tracks", "conditions", "difficult_examples"}
    }
    parts.append(
        "<details><summary>Frozen configuration, environment and evidence</summary>"
        f"<pre>{_escape(json.dumps(context, ensure_ascii=False, indent=2))}</pre>"
        "</details></section>"
    )
    examples = report.get("difficult_examples", [])
    if examples:
        parts.append("<section><h2>Difficult examples</h2>")
        for example in examples:
            parts.append(
                f"<h3>{_escape(example['model'])} / {_escape(example['sample_id'])}"
                f" · pass {_escape(example.get('pass'))}</h3>"
            )
            if example.get("image"):
                parts.append(
                    f'<img alt="Frozen handwriting sample" src="{_escape(example["image"])}">'
                )
            parts.append(
                _table(
                    [
                        {
                            key: example.get(key)
                            for key in ("status", "reference", "prediction", "cer", "edit_counts")
                        }
                    ]
                )
            )
        parts.append("</section>")
    parts.append("</main></html>")
    return "".join(parts)


def _evidence(report, source, target):
    """Attach bounded difficult examples and operational context to the shared model."""
    from .inspection import inspect_run
    from .runner import read_records

    source = Path(source).resolve()
    if (source / "manifest.json").is_file():
        paths = [source]
    elif (source / "experiment.json").is_file():
        paths = json.loads((source / "experiment.json").read_text())["runs"]
    elif (source / "suite.json").is_file():
        state = json.loads((source / "suite.json").read_text())
        paths = [entry["run_dir"] for entry in state["runs"].values() if entry.get("run_dir")]
    else:
        paths = []
    examples, contexts, operations = [], [], []
    timings = defaultdict(list)
    for value in paths:
        directory = Path(value)
        if not directory.is_absolute():
            directory = source / directory
        if not directory.is_dir():
            directory = source / Path(value).name
        manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
        context = {
            key: manifest.get(key)
            for key in (
                "benchmark_fingerprint",
                "repetition",
                "suite_condition",
                "model_info",
                "host",
                "provider_controls",
                "measurement_profile",
                "settings",
                "prompt_hashes",
            )
        }
        execution = directory / "execution.json"
        context["execution"] = json.loads(execution.read_text()) if execution.is_file() else None
        context["cost_status"] = "unknown unless explicitly estimated in execution evidence"
        contexts.append(context)
        records = read_records(directory / "results.jsonl")
        attempts = read_records(directory / "attempts.jsonl")
        state = context["execution"] or {}
        invocations = state.get("invocations", [])
        wall = sum(value.get("workload_wall_seconds") or 0 for value in invocations)
        successful = sum(
            row.get("status") == "success" and not row.get("cache_hit") for row in records
        )
        starts = sum(row.get("event_kind") == "request_started" for row in attempts)
        finishes = sum(row.get("event_kind") == "request_finished" for row in attempts)
        cost_settings = manifest.get("costs") or {}
        priced = any(
            isinstance(cost_settings.get(key), dict)
            and any(
                name in cost_settings[key]
                for name in (
                    "cost_per_sample_usd",
                    "operating_cost_per_sample_usd",
                    "input_per_million_usd",
                )
            )
            for key in ("local", "api")
        )
        operations.append(
            {
                "pass": manifest.get("repetition", 1),
                "status": manifest.get("status"),
                "requests": state.get("generation_requests_cumulative"),
                "retryable_failures": sum(
                    row.get("status") == "retryable_error" for row in attempts
                ),
                "uncertain_attempts": max(0, starts - finishes),
                "failed_predictions": sum(row.get("status") == "error" for row in records),
                "workload_wall_seconds": wall or None,
                "successful_samples_per_second": successful / wall if wall else None,
                "throughput_denominator": "elapsed workload wall seconds, including retries",
                "estimated_cost_usd": state.get("estimated_spend_usd_cumulative")
                if priced
                else None,
                "cost_status": "estimated"
                if priced and not state.get("unknown_spend_count_cumulative")
                else "unknown",
            }
        )
        for row in records:
            if row.get("cache_hit") or row.get("status") != "success":
                continue
            for timing in row.get("timing_observations", []):
                value = timing.get("value")
                if (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and math.isfinite(value)
                ):
                    key = (
                        row["model"],
                        timing["metric_name"],
                        timing["unit"],
                        timing["source"],
                        timing.get("start_boundary"),
                        timing.get("end_boundary"),
                    )
                    timings[key].append(value)
        if len(examples) >= 12:
            continue
        for row in inspect_run(directory, worst=min(4, 12 - len(examples)))["results"]:
            example = {
                key: row.get(key)
                for key in (
                    "model",
                    "sample_id",
                    "status",
                    "reference",
                    "prediction",
                    "cer",
                    "flags",
                )
            }
            example["edit_counts"] = row["alignment"].get("counts")
            example["pass"] = manifest.get("repetition") or 1
            example["condition"] = manifest.get("suite_condition") or "baseline"
            example["image"] = None
            sample = row.get("sample", {})
            image = Path(sample.get("crop_path") or "")
            if not image.is_absolute():
                image = directory / image
            if not image.is_file():
                old_parts = Path(sample.get("crop_path", "")).parts
                markers = [index for index, part in enumerate(old_parts) if part == "crops"]
                image = (
                    directory
                    / "crops"
                    / (Path(*old_parts[markers[-1] + 1 :]) if markers else Path(*old_parts[-1:]))
                )
            try:
                image.resolve().relative_to((directory / "crops").resolve())
            except ValueError:
                image = None
            if image and image.is_file() and image.suffix.lower() in {".png", ".jpg", ".jpeg"}:
                content_hash = hashlib.sha256(image.read_bytes()).hexdigest()
                expected = sample.get("hashes", {}).get("crop")
                if expected and content_hash != expected:
                    raise ValueError(f"Frozen image hash mismatch for {row['sample_id']!r}")
                assets = target / "assets"
                assets.mkdir(exist_ok=True)
                name = f"{content_hash}{image.suffix.lower()}"
                shutil.copyfile(image, assets / name)
                example["image"] = f"assets/{name}"
            examples.append(example)
    report["measurement_context"] = contexts
    report["difficult_examples"] = examples
    report["operational_summary"] = operations
    report["timing_breakdowns"] = [
        dict(
            zip(("model", "metric", "unit", "source", "start_boundary", "end_boundary"), key),
            median=statistics.median(values),
            observations=len(values),
        )
        for key, values in sorted(timings.items(), key=lambda item: str(item[0]))
    ]


def write_report(source, output=None, *, formats=("html", "xlsx", "csv", "json"), policy=None):
    """Evaluate saved evidence once; render every format from the exact same model."""
    from .report_model import build_report

    formats = tuple(formats)
    invalid = set(formats) - {"html", "xlsx", "csv", "json"}
    if invalid:
        raise ValueError(f"Unsupported report formats: {sorted(invalid)}")
    report = build_report(Path(source), policy=policy)
    target = Path(output) if output else Path(source) / "report"
    target.mkdir(parents=True, exist_ok=True)
    _evidence(report, source, target)
    paths = []
    rows = [_flat(row) for row in _leaderboards(report)]
    fields = list(dict.fromkeys(key for row in rows for key in row)) or [
        "condition",
        "track",
        "model",
    ]
    if "json" in formats:
        path = target / "evaluation.json"
        _atomic_json(path, report)
        paths.append(path)
    if "html" in formats:
        path = target / "report.html"
        path.write_text(_html(report), encoding="utf-8")
        paths.append(path)
    if "csv" in formats:
        path = target / "leaderboards.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows({key: _csv_value(row.get(key)) for key in fields} for row in rows)
        paths.append(path)
    if "xlsx" in formats:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leaderboards"
        overflow = []
        summary_fields = [
            key
            for key in (
                "condition",
                "track",
                "rank",
                "model",
                "score",
                "score_stddev",
                "raw_accuracy_rank",
                "corpus_cer",
                "wer",
                "exact_match_rate",
                "median_latency_seconds",
                "p95_latency_seconds",
                "status",
            )
            if key in fields
        ]
        _append_table(sheet, sheet.title, summary_fields, rows, overflow)
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="007B78")
            cell.font = Font(color="FFFFFF", bold=True)
        for index, field in enumerate(summary_fields, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                52, max(18, len(field) + 3)
            )
        tables = {
            "Operations": report.get("operational_summary", []),
            "Timing": report.get("timing_breakdowns", []),
            "Passes": [
                {"condition": condition, "track": track, "model": model["model"], **item}
                for condition, evaluation in _evaluations(report)
                for track, data in evaluation.get("tracks", {}).items()
                for model in data.get("models", [])
                for item in model.get("per_pass", [])
            ],
        }
        for title, table in tables.items():
            if table:
                columns = list(dict.fromkeys(key for row in table for key in row))
                _append_table(workbook.create_sheet(title), title, columns, table, overflow)
        # Entire report remains available in Excel, including null/unknown provenance.
        evidence = workbook.create_sheet("Evaluation model")
        evidence_rows = []

        def walk(value, prefix=""):
            if isinstance(value, dict):
                for key, child in value.items():
                    walk(child, f"{prefix}.{key}" if prefix else str(key))
            elif isinstance(value, list):
                for index, child in enumerate(value):
                    walk(child, f"{prefix}[{index}]")
            else:
                evidence_rows.append({"JSON path": prefix, "Value": _text(value)})

        walk(report)
        _append_table(evidence, evidence.title, ["JSON path", "Value"], evidence_rows, overflow)
        evidence.freeze_panes = "B2"
        evidence.column_dimensions["A"].width = 70
        evidence.column_dimensions["B"].width = 90
        if overflow:
            chunks = workbook.create_sheet("Text Overflow")
            chunks.append(["ID", "Sheet", "Row", "Column", "Chunk", "Chunks", "Text"])
            for row in overflow:
                chunks.append(row)
                for cell in chunks[chunks.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        path = target / "report.xlsx"
        workbook.save(path)
        paths.append(path)
    return paths
