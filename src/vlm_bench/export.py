"""Export canonical benchmark results as CSV and Excel summaries."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable

from .metrics import summarize

_SAMPLE_FIELD_ORDER = [
    "model",
    "sample_id",
    "status",
    "reference",
    "prediction",
    "cer",
    "wer",
    "raw_cer",
    "exact_match",
    "char_substitutions",
    "char_deletions",
    "char_insertions",
    "char_edits",
    "word_edits",
    "reference_chars",
    "reference_words",
    "latency_seconds",
    "load_duration_seconds",
    "truncated",
    "error",
]
_SUMMARY_FIELD_ORDER = [
    "rank",
    "model",
    "complete",
    "sample_count",
    "expected_sample_count",
    "missing_sample_count",
    "sample_coverage_complete",
    "expected_model",
    "scored_sample_count",
    "success_count",
    "failed_count",
    "failure_rate",
    "empty_count",
    "truncated_count",
    "cer",
    "wer",
    "mean_sample_cer",
    "mean_sample_wer",
    "exact_match_rate",
    "mean_latency_seconds",
    "median_latency_seconds",
    "p95_latency_seconds",
    "total_latency_seconds",
    "samples_per_minute",
    "load_duration_seconds",
    "load_event_count",
    "warmup_count",
    "warmup_load_duration_seconds",
]
_CSV_FORMULA_PREFIX = re.compile(r"^[\ufeff\s]*[=+\-@]")
_XML_ILLEGAL_CONTROL = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_EXCEL_JSON_STRING_PREFIX = "[VLM-BENCH:JSON-STRING-V1]"
_EXCEL_TEXT_LIMIT_UTF16 = 32_760


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _flatten_record(record: dict[str, Any]) -> dict[str, Any]:
    """Put nested metric fields beside their sample fields for tabular output."""
    flattened: dict[str, Any] = {}
    for key, value in record.items():
        if key == "metrics" and isinstance(value, dict):
            for metric_name, metric_value in value.items():
                flattened.setdefault(metric_name, metric_value)
        elif key != "metrics":
            flattened[key] = value
    return flattened


def _table_columns(rows: list[dict[str, Any]], preferred: Iterable[str] = ()) -> list[str]:
    present = {key for row in rows for key in row}
    ordered = [key for key in preferred if key in present or not rows]
    extras = sorted(present.difference(ordered))
    return ordered + extras


def _csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list, tuple)):
        value = _json_text(value)
    if isinstance(value, str) and _CSV_FORMULA_PREFIX.match(value):
        # Prefixing the cell with an apostrophe prevents spreadsheet apps from
        # evaluating formula-like text when a CSV is opened interactively.
        return "'" + value
    return value


def _write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key)) for key in columns})


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dict, list, tuple)):
        return _json_text(value)
    value = str(value)
    if _XML_ILLEGAL_CONTROL.search(value):
        # XML 1.0 cannot store these code points. Keep their exact value in a
        # tagged JSON string so it can be restored by stripping this prefix
        # and JSON-decoding the remainder. Canonical JSONL remains untouched.
        return _EXCEL_JSON_STRING_PREFIX + json.dumps(value, ensure_ascii=False)
    return value


def _utf16_length(text: str) -> int:
    return len(text.encode("utf-16-le")) // 2


def _chunks(text: str, maximum_units: int = _EXCEL_TEXT_LIMIT_UTF16) -> list[str]:
    """Split text below Excel's per-cell limit without cutting surrogate pairs."""
    result: list[str] = []
    current: list[str] = []
    units = 0
    for character in text:
        char_units = _utf16_length(character)
        if current and units + char_units > maximum_units:
            result.append("".join(current))
            current = []
            units = 0
        current.append(character)
        units += char_units
    if current or not result:
        result.append("".join(current))
    return result


def _append_table(
    worksheet: Any,
    title: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    overflow_rows: list[list[Any]],
) -> None:
    worksheet.append(columns)
    for row_number, row in enumerate(rows, start=2):
        values: list[Any] = []
        for column_number, column in enumerate(columns, start=1):
            value = _excel_value(row.get(column))
            if isinstance(value, str) and _utf16_length(value) > _EXCEL_TEXT_LIMIT_UTF16:
                overflow_id = f"OVF{len(overflow_rows) + 1:07d}"
                text_chunks = _chunks(value)
                marker = f"[FULL TEXT IN Text Overflow: {overflow_id}]"
                overflow_rows.extend(
                    [
                        overflow_id,
                        title,
                        row_number,
                        column,
                        chunk_index,
                        len(text_chunks),
                        chunk,
                    ]
                    for chunk_index, chunk in enumerate(text_chunks, start=1)
                )
                value = marker
            values.append(value)
        worksheet.append(values)
        for cell in worksheet[row_number]:
            # openpyxl treats strings beginning with '=' as formulas. Force all
            # textual values to stored strings, including formula-like content.
            if isinstance(cell.value, str):
                cell.data_type = "s"

    worksheet.freeze_panes = "A2"
    if columns and rows:
        worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 24
    for cell in worksheet[1]:
        from openpyxl.styles import Font, PatternFill

        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="2F5597")
    for index, column in enumerate(columns, start=1):
        header = worksheet.cell(row=1, column=index).value
        width = min(max(len(str(header)) + 2, 12), 48)
        worksheet.column_dimensions[worksheet.cell(row=1, column=index).column_letter].width = width
        if str(header).lower() in {"prediction", "reference", "error", "text", "value"}:
            from openpyxl.styles import Alignment

            for cells in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for cell in cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    percent_columns = {
        "cer",
        "wer",
        "raw_cer",
        "mean_sample_cer",
        "mean_sample_wer",
        "exact_match_rate",
    }
    for index, column in enumerate(columns, start=1):
        if column in percent_columns:
            for row_number in range(2, len(rows) + 2):
                worksheet.cell(row=row_number, column=index).number_format = "0.00%"


def _flatten_manifest(manifest: dict[str, Any]) -> list[dict[str, str]]:
    flattened: list[dict[str, str]] = []

    def visit(prefix: str, value: Any) -> None:
        if isinstance(value, dict) and value:
            for key in sorted(value):
                child_key = f"{prefix}.{key}" if prefix else str(key)
                visit(child_key, value[key])
        else:
            flattened.append(
                {
                    "key": prefix or "manifest",
                    "value": _json_text(value) if isinstance(value, (dict, list)) else str(value),
                }
            )

    for key in sorted(manifest):
        visit(str(key), manifest[key])
    return flattened


def _write_workbook(
    path: Path,
    summaries: list[dict[str, Any]],
    sample_rows: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    manifest: dict[str, Any],
    warmup_rows: list[dict[str, Any]],
) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - depends on installation
        raise RuntimeError(
            "Excel export requires openpyxl; install the export dependencies"
        ) from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    overflow_rows: list[list[Any]] = []
    tables = [
        ("Summary", summaries, _table_columns(summaries, _SUMMARY_FIELD_ORDER)),
        ("Samples", sample_rows, _table_columns(sample_rows, _SAMPLE_FIELD_ORDER)),
        ("Errors", errors, _table_columns(errors, _SAMPLE_FIELD_ORDER)),
        (
            "Run Configuration",
            _flatten_manifest(manifest),
            ["key", "value"],
        ),
    ]
    if warmup_rows:
        flattened_warmups = [_flatten_record(record) for record in warmup_rows]
        tables.append(
            (
                "Warmups",
                flattened_warmups,
                _table_columns(
                    flattened_warmups,
                    [
                        "model",
                        "sample_id",
                        "status",
                        "latency_seconds",
                        "load_duration_seconds",
                        "error",
                    ],
                ),
            )
        )
    for title, rows, columns in tables:
        worksheet = workbook.create_sheet(title)
        _append_table(worksheet, title, columns, rows, overflow_rows)

    if overflow_rows:
        worksheet = workbook.create_sheet("Text Overflow")
        overflow_headers = [
            "overflow_id",
            "sheet",
            "source_row",
            "field",
            "chunk_index",
            "chunk_count",
            "text_chunk",
        ]
        _append_table(
            worksheet,
            "Text Overflow",
            overflow_headers,
            [dict(zip(overflow_headers, row)) for row in overflow_rows],
            [],
        )
        worksheet.column_dimensions["G"].width = 90
        from openpyxl.styles import Alignment

        for cells in worksheet.iter_cols(min_col=7, max_col=7, min_row=2):
            for cell in cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(path)


def _read_results(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as file:
        for line_number, line in enumerate(file, start=1):
            if not line.strip():
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"invalid JSON in {path.name} line {line_number}: {exc}") from exc
            if not isinstance(record, dict):
                raise ValueError(f"{path.name} line {line_number} must contain a JSON object")
            records.append(record)
    return records


def export_run(run_dir: Path, formats: list[str]) -> list[Path]:
    """Export an existing run without modifying canonical ``results.jsonl``.

    Supported formats are ``xlsx``, ``csv``, and ``jsonl``. CSV writes separate
    ``summary.csv`` and ``samples.csv`` files; Excel writes the summary, samples,
    failures, and manifest configuration to one workbook. Text too long for an
    Excel cell is stored in ordered chunks on a ``Text Overflow`` sheet, and
    the original cell contains an explicit overflow ID marker. XML-illegal
    control characters are stored using a tagged JSON string representation.
    """
    normalized_formats = list(dict.fromkeys(str(value).lower().lstrip(".") for value in formats))
    unsupported = [value for value in normalized_formats if value not in {"xlsx", "csv", "jsonl"}]
    if unsupported:
        raise ValueError(f"unsupported export format(s): {', '.join(unsupported)}")

    run_dir = Path(run_dir)
    results_path = run_dir / "results.jsonl"
    manifest_path = run_dir / "manifest.json"
    records = _read_results(results_path)
    warmups_path = run_dir / "warmups.jsonl"
    warmups = _read_results(warmups_path) if warmups_path.exists() else []
    with manifest_path.open("r", encoding="utf-8") as file:
        manifest = json.load(file)
    if not isinstance(manifest, dict):
        raise ValueError("manifest.json must contain a JSON object")

    sample_rows = [_flatten_record(record) for record in records]
    expected_samples = manifest.get("samples")
    if isinstance(expected_samples, list):
        expected_samples = [
            sample.get("id")
            for sample in expected_samples
            if isinstance(sample, dict) and "id" in sample
        ]
    elif not isinstance(expected_samples, int):
        expected_samples = None
    expected_models = manifest.get("models")
    if not isinstance(expected_models, list):
        expected_models = None
    summary_rows = summarize(
        records,
        expected_samples=expected_samples,
        models=expected_models,
        force_incomplete=manifest.get("status") in {"running", "interrupted"},
        warmups=warmups,
    )
    summary_columns = _table_columns(summary_rows, _SUMMARY_FIELD_ORDER)
    sample_columns = _table_columns(sample_rows, _SAMPLE_FIELD_ORDER)
    error_rows = [
        row
        for row, record in zip(sample_rows, records)
        if record.get("status") != "success" or record.get("error")
    ]
    outputs: list[Path] = []

    if "xlsx" in normalized_formats:
        workbook_path = run_dir / "results.xlsx"
        _write_workbook(workbook_path, summary_rows, sample_rows, error_rows, manifest, warmups)
        outputs.append(workbook_path)
    if "csv" in normalized_formats:
        summary_path = run_dir / "summary.csv"
        samples_path = run_dir / "samples.csv"
        _write_csv(summary_path, summary_rows, summary_columns)
        _write_csv(samples_path, sample_rows, sample_columns)
        outputs.extend([summary_path, samples_path])
    if "jsonl" in normalized_formats:
        outputs.append(results_path)
    return outputs
